"""Word documents: architecture decision records.

A .docx is a zip archive of XML. python-docx handles the unpacking; we join
the paragraphs it finds, dropping empties, and read the title and author from
the file's own core properties where present.
"""

from __future__ import annotations

from pathlib import Path

import docx

from ragsys.ingestion.loaders._paths import citation_path, document_id
from ragsys.ingestion.models import RawDocument


def load_docx(path: Path) -> list[RawDocument]:
    document = docx.Document(str(path))
    paragraphs = [p.text.strip() for p in document.paragraphs if p.text.strip()]
    if not paragraphs:
        return []

    core = document.core_properties
    return [
        RawDocument(
            doc_id=document_id(path),
            text="\n\n".join(paragraphs),
            source_path=citation_path(path),
            fmt="docx",
            metadata={
                "title": (core.title or paragraphs[0])[:200],
                "author": core.author or None,
            },
        )
    ]
