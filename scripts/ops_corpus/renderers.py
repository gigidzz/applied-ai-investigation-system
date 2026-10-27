"""Render the structured facts into the six document formats.

Each function takes one fact object and writes one file. Keeping rendering
separate from the facts means a format can be changed without touching the
content, and the same fact can appear in two formats — which is how the
cross-format near-duplicate case is produced.

Both binary formats are made reproducible on purpose: PDFs and DOCX files
normally embed a creation timestamp, which would make every regeneration
produce different bytes and break the "repeatable evaluation" requirement.
"""

from __future__ import annotations

import csv
import json
import zipfile
from datetime import UTC, datetime
from pathlib import Path
from xml.sax.saxutils import escape

import docx
import reportlab.rl_config
from docx.shared import Pt
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

from .facts import (
    COMPANY,
    Adr,
    Incident,
    OpsReport,
    Runbook,
    Service,
    Thread,
)

# Strip the creation timestamp and document id from generated PDFs so that
# regenerating the corpus produces identical bytes.
reportlab.rl_config.invariant = 1

FIXED_TIMESTAMP = datetime(2026, 9, 1, 12, 0, 0, tzinfo=UTC)


# --------------------------------------------------------------------------- #
# PDF — incident post-mortems
# --------------------------------------------------------------------------- #


def _pdf_styles() -> dict[str, ParagraphStyle]:
    base = getSampleStyleSheet()
    return {
        "title": ParagraphStyle(
            "pmTitle", parent=base["Title"], fontSize=16, spaceAfter=14, alignment=0
        ),
        "heading": ParagraphStyle(
            "pmHeading", parent=base["Heading2"], fontSize=11.5, spaceBefore=12, spaceAfter=5
        ),
        "body": ParagraphStyle(
            "pmBody", parent=base["BodyText"], fontSize=9.5, leading=13.5, spaceAfter=5
        ),
        "meta": ParagraphStyle(
            "pmMeta", parent=base["BodyText"], fontSize=8.5, leading=12, textColor="#444444"
        ),
    }


def write_incident_pdf(incident: Incident, path: Path) -> None:
    """Write one incident post-mortem as a PDF."""
    styles = _pdf_styles()
    flow = [
        Paragraph(escape(f"{incident.incident_id}: {incident.title}"), styles["title"]),
        Paragraph(
            escape(
                f"{COMPANY} · {incident.severity} · service {incident.service} · "
                f"author {incident.author}"
            ),
            styles["meta"],
        ),
        Paragraph(
            escape(
                f"Started {incident.started_at} · detected {incident.detected_at} · "
                f"resolved {incident.resolved_at}"
            ),
            styles["meta"],
        ),
        Spacer(1, 10),
    ]

    sections: list[tuple[str, str]] = [
        ("Customer impact", incident.impact),
        ("Root cause", incident.root_cause),
    ]
    if incident.caused_by_release:
        sections.append(
            (
                "Causing change",
                f"This incident was caused by release {incident.caused_by_release}, "
                f"deployed to {incident.service}. The release is recorded in the "
                f"change log.",
            )
        )
    sections.append(
        (
            "Affected services",
            ", ".join(incident.affected_services)
            + ". Downstream services were affected through their dependency on "
            + incident.service
            + ".",
        )
    )

    for heading, body in sections:
        flow.append(Paragraph(escape(heading), styles["heading"]))
        flow.append(Paragraph(escape(body), styles["body"]))

    flow.append(Paragraph("Timeline (UTC)", styles["heading"]))
    for entry in incident.timeline:
        flow.append(Paragraph(escape(f"{entry.at} — {entry.note}"), styles["body"]))

    flow.append(Paragraph("Action items", styles["heading"]))
    for item in incident.action_items:
        flow.append(Paragraph(escape(f"• {item}"), styles["body"]))

    doc = SimpleDocTemplate(
        str(path),
        pagesize=A4,
        leftMargin=2.2 * cm,
        rightMargin=2.2 * cm,
        topMargin=2 * cm,
        bottomMargin=2 * cm,
        title=f"{incident.incident_id} {incident.title}",
        author=incident.author,
        subject=f"Post-mortem for {incident.service}",
    )
    doc.build(flow)


# --------------------------------------------------------------------------- #
# CSV — the change log
# --------------------------------------------------------------------------- #

CHANGE_LOG_COLUMNS = (
    "release",
    "service",
    "deployed_at",
    "author",
    "risk",
    "rolled_back",
    "summary",
)


def write_change_log_csv(deploys: tuple, path: Path) -> None:
    """Write every deploy as one row. Each row is an independent record."""
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(CHANGE_LOG_COLUMNS)
        for deploy in deploys:
            writer.writerow(
                [
                    deploy.release,
                    deploy.service,
                    deploy.deployed_at,
                    deploy.author,
                    deploy.risk,
                    "yes" if deploy.rolled_back else "no",
                    deploy.summary,
                ]
            )


# --------------------------------------------------------------------------- #
# JSON — the service catalogue
# --------------------------------------------------------------------------- #

# The catalogue records settlement-queue as owned by Payments Core. ADR-0007
# transfers it to Platform, which is contradiction C2: the catalogue is stale.
SHARED_INFRASTRUCTURE = {
    "settlement-queue": {"owner_team": "Payments Core", "kind": "message queue"},
    "redis-cluster": {"owner_team": "Platform", "kind": "cache"},
    "postgres-primary": {"owner_team": "Payments Core", "kind": "database"},
    "postgres-replica": {"owner_team": "Payments Core", "kind": "database"},
    "feature-store": {"owner_team": "Risk", "kind": "service"},
    "document-store": {"owner_team": "Compliance", "kind": "object storage"},
    "email-relay": {"owner_team": "Platform", "kind": "third party"},
}


def write_service_json(service: Service, all_services: tuple[Service, ...], path: Path) -> None:
    """Write one service catalogue entry, including who depends on it."""
    record = {
        "name": service.name,
        "team": service.team,
        "tier": service.tier,
        "summary": service.summary,
        "repository": service.repo,
        "on_call_rotation": service.on_call,
        "depends_on": list(service.depends_on),
        "depended_on_by": [
            other.name for other in all_services if service.name in other.depends_on
        ],
        "shared_infrastructure": {
            dep: SHARED_INFRASTRUCTURE[dep]
            for dep in service.depends_on
            if dep in SHARED_INFRASTRUCTURE
        },
        "catalogue_version": "2026-01-10",
    }
    path.write_text(json.dumps(record, indent=2) + "\n", encoding="utf-8")


# --------------------------------------------------------------------------- #
# DOCX — architecture decision records
# --------------------------------------------------------------------------- #


def _normalise_docx(path: Path) -> None:
    """Rewrite a .docx zip with fixed member timestamps.

    A .docx is a zip archive, and python-docx stamps every member with the
    current wall-clock time. Two runs therefore produce different bytes for
    identical content, which would invalidate the corpus fingerprint and, with
    it, the claim that evaluation is repeatable. Rewriting the archive with a
    fixed timestamp and a normalised host field makes regeneration
    byte-identical.
    """
    fixed = (
        FIXED_TIMESTAMP.year,
        FIXED_TIMESTAMP.month,
        FIXED_TIMESTAMP.day,
        FIXED_TIMESTAMP.hour,
        FIXED_TIMESTAMP.minute,
        FIXED_TIMESTAMP.second,
    )
    with zipfile.ZipFile(path) as source:
        members = [(info, source.read(info.filename)) for info in source.infolist()]

    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as target:
        for info, payload in members:
            normalised = zipfile.ZipInfo(info.filename, date_time=fixed)
            normalised.compress_type = info.compress_type
            normalised.external_attr = info.external_attr
            normalised.internal_attr = info.internal_attr
            normalised.create_system = 0
            target.writestr(normalised, payload)


def write_adr_docx(adr: Adr, path: Path) -> None:
    document = docx.Document()

    style = document.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)

    document.add_heading(f"{adr.adr_id}: {adr.title}", level=1)

    meta = document.add_paragraph()
    meta.add_run(f"Status: {adr.status}    Date: {adr.date}    Author: {adr.author}").italic = True

    document.add_heading("Context", level=2)
    document.add_paragraph(adr.context)

    document.add_heading("Decision", level=2)
    document.add_paragraph(adr.decision)

    document.add_heading("Consequences", level=2)
    for consequence in adr.consequences:
        document.add_paragraph(consequence, style="List Bullet")

    core = document.core_properties
    core.author = adr.author
    core.title = f"{adr.adr_id} {adr.title}"
    core.created = FIXED_TIMESTAMP
    core.modified = FIXED_TIMESTAMP
    core.revision = 1

    document.save(str(path))
    _normalise_docx(path)


# --------------------------------------------------------------------------- #
# Markdown — runbooks, reports, summaries
# --------------------------------------------------------------------------- #


def write_runbook_md(runbook: Runbook, path: Path) -> None:
    lines = [f"# {runbook.title}", ""]
    if runbook.service:
        lines += [f"**Service:** `{runbook.service}`", ""]
    for heading, body in runbook.sections:
        lines += [f"## {heading}", "", body, ""]
    path.write_text("\n".join(lines).rstrip() + "\n", encoding="utf-8")


def write_ops_report_md(report: OpsReport, path: Path) -> None:
    lines = [
        f"# Operations report — {report.period}",
        "",
        f"**Deploys:** {report.deploy_count}",
        "",
        "## Highlights",
        "",
    ]
    lines += [f"- {item}" for item in report.highlights]
    lines += ["", "## Incidents", ""]
    if report.incident_ids:
        lines += [f"- {incident_id}" for incident_id in report.incident_ids]
    else:
        lines += ["- None filed this period."]
    path.write_text("\n".join(lines).rstrip() + "\n", encoding="utf-8")


def write_incident_summary_md(incident: Incident, path: Path) -> None:
    """A condensed Markdown restatement of a post-mortem.

    This exists so deduplication has a cross-format near-duplicate to catch:
    most of the wording is shared with the PDF.
    """
    lines = [
        f"# Summary: {incident.incident_id}",
        "",
        f"**Service:** `{incident.service}` · **Severity:** {incident.severity}",
        "",
        "## What happened",
        "",
        incident.impact,
        "",
        "## Why",
        "",
        incident.root_cause,
        "",
        "## Resolution",
        "",
        f"Resolved at {incident.resolved_at}. "
        + (
            f"Caused by release {incident.caused_by_release}."
            if incident.caused_by_release
            else "No causing release identified."
        ),
    ]
    path.write_text("\n".join(lines).rstrip() + "\n", encoding="utf-8")


# --------------------------------------------------------------------------- #
# TXT — chat threads
# --------------------------------------------------------------------------- #


def write_thread_txt(thread: Thread, path: Path) -> None:
    lines = [
        f"channel: {thread.channel}",
        f"date: {thread.date}",
        f"thread: {thread.thread_id}",
        "-" * 60,
    ]
    for message in thread.messages:
        lines.append(f"[{message.at}] {message.who}: {message.text}")
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")
